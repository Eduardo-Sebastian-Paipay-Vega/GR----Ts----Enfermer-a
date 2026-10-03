import docx
from docx.shared import Pt, Cm

doc = docx.Document(r"c:\GRESLY\DOCUMENTOS\Generación\protocolo_cualitativo.docx")
print("=== AUDITORÍA DOCUMENTO WORD ===")
print("Total secciones:", len(doc.sections))
for i, s in enumerate(doc.sections):
    h = " ".join(p.text for p in s.header.paragraphs if p.text.strip())
    f = " ".join(p.text for p in s.footer.paragraphs if p.text.strip())
    print(f"Sec {i}: Margins (T={s.top_margin.cm:.1f}, L={s.left_margin.cm:.1f}, R={s.right_margin.cm:.1f}, B={s.bottom_margin.cm:.1f}) | Header='{h}' | Footer='{f}'")

print("\nPrimeros 25 párrafos:")
for i, p in enumerate(doc.paragraphs[:25]):
    if p.text.strip():
        print(f"P{i:02d} [Indent={p.paragraph_format.first_line_indent.cm if p.paragraph_format.first_line_indent else 0:.2f}cm, Align={p.alignment}]: {p.text[:80]}")

print("\nÚltimos 15 párrafos (Referencias):")
for i, p in enumerate(doc.paragraphs[-15:]):
    if p.text.strip():
        print(f"Ref {i:02d} [LeftIndent={p.paragraph_format.left_indent.cm if p.paragraph_format.left_indent else 0:.2f}cm]: {p.text[:90]}")
